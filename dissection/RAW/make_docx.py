"""
Convert UMS_REPORT.md and LPU_TOUCH_REPORT.md to professional .docx files.
Uses python-docx with proper heading hierarchy, tables, code blocks, and styling.
"""
import re
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DISSECTION = '/mnt/c/Saumya_workspace/codeflux/dissection'

# ─────────────────────────────────────────────────────────────
# THEME COLORS
# ─────────────────────────────────────────────────────────────
COLOR_TITLE      = RGBColor(0x1A, 0x1A, 0x2E)  # Deep navy
COLOR_H1         = RGBColor(0x16, 0x21, 0x3E)  # Dark blue
COLOR_H2         = RGBColor(0x0F, 0x3D, 0x6B)  # Medium blue
COLOR_H3         = RGBColor(0x1E, 0x5F, 0x99)  # Lighter blue
COLOR_ACCENT     = RGBColor(0x6C, 0x63, 0xFF)  # Electric violet (brand)
COLOR_CODE_BG    = RGBColor(0xF4, 0xF4, 0xF5)  # Light grey
COLOR_TABLE_HDR  = RGBColor(0x1A, 0x1A, 0x2E)  # Same as title
COLOR_TABLE_ALT  = RGBColor(0xF8, 0xF8, 0xFF)  # Ghost white
COLOR_BODY       = RGBColor(0x1A, 0x1A, 0x2E)  # Near black


def set_cell_bg(cell, color_hex):
    """Set table cell background color."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def set_para_border_bottom(para, color='6C63FF', size=6):
    """Add a bottom border to a paragraph (used under H1)."""
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_horizontal_rule(doc):
    """Add a thin horizontal rule paragraph."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CCCCCC')
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)


def style_doc(doc):
    """Apply base document styles."""
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = COLOR_BODY

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)


def add_title(doc, text, subtitle=None):
    """Add document title block."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = COLOR_TITLE
    run.font.name = 'Calibri'
    set_para_border_bottom(p, 'AAAAAA', 8)

    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(10)
        r2 = p2.add_run(subtitle)
        r2.font.size = Pt(13)
        r2.font.color.rgb = COLOR_H2
        r2.font.name = 'Calibri'
        r2.font.bold = False


def add_heading(doc, text, level):
    """Add styled heading at level 1/2/3."""
    p = doc.add_paragraph()
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text.upper())
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_H1
        run.font.name = 'Calibri'
        set_para_border_bottom(p, '6C63FF', 6)
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = COLOR_H2
        run.font.name = 'Calibri'
    elif level == 3:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = COLOR_H3
        run.font.name = 'Calibri'


def add_blockquote(doc, text):
    """Add a styled blockquote paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    # Left border
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '12')
    left.set(qn('w:space'), '6')
    left.set(qn('w:color'), '6C63FF')
    pBdr.append(left)
    pPr.append(pBdr)
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x66)
    run.font.name = 'Calibri'


def add_code_block(doc, text):
    """Add a monospace code block with grey background."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    # Shade background
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F0F0F5')
    pPr.append(shd)
    run = p.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)


def add_body_para(doc, text):
    """Add a styled body paragraph with inline bold/code handling."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    _add_inline_runs(p, text)


def add_bullet(doc, text, level=0):
    """Add a bullet list item."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.2)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    _add_inline_runs(p, text)


def _add_inline_runs(p, text):
    """Parse **bold**, `code`, and plain text inline."""
    pattern = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`)')
    parts = pattern.split(text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(10.5)
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9.5)
            run.font.color.rgb = COLOR_ACCENT
        else:
            run = p.add_run(part)
            run.font.name = 'Calibri'
            run.font.size = Pt(10.5)


def add_table_from_md(doc, lines):
    """Parse and render a markdown table."""
    rows = []
    for line in lines:
        if re.match(r'^\s*\|[-:| ]+\|\s*$', line):
            continue
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        rows.append(cells)

    if not rows:
        return

    max_cols = max(len(r) for r in rows)
    # Normalize row lengths
    rows = [r + [''] * (max_cols - len(r)) for r in rows]

    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)

            # Strip markdown bold from cell text
            clean = re.sub(r'\*\*([^*]+)\*\*', r'\1', cell_text)

            if i == 0:
                # Header row
                set_cell_bg(cell, '1A1A2E')
                run = p.add_run(clean)
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(9.5)
                run.font.name = 'Calibri'
            else:
                # Alternating row color
                if i % 2 == 0:
                    set_cell_bg(cell, 'F0F0FF')
                run = p.add_run(clean)
                run.font.size = Pt(9.5)
                run.font.name = 'Calibri'

    doc.add_paragraph()  # spacing after table


def add_meta_box(doc, lines):
    """Render the > blockquote lines as a styled info box."""
    text = ' '.join(
        re.sub(r'^>\s*\*?\*?', '', l).replace('**', '').strip()
        for l in lines
    )
    add_blockquote(doc, text)


def md_to_docx(md_path, docx_path, doc_title, doc_subtitle):
    """Main converter: read markdown, write docx."""
    doc = Document()
    style_doc(doc)

    with open(md_path, 'r', encoding='utf-8') as f:
        raw = f.read()

    lines = raw.split('\n')

    add_title(doc, doc_title, doc_subtitle)

    i = 0
    table_buffer = []
    code_buffer = []
    quote_buffer = []
    in_code = False
    in_table = False

    while i < len(lines):
        line = lines[i]

        # ── Code block (triple backtick) ──────────────────────
        if line.strip().startswith('```'):
            if in_code:
                # End code block
                add_code_block(doc, '\n'.join(code_buffer))
                code_buffer = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        # ── Table detection ────────────────────────────────────
        if '|' in line and line.strip().startswith('|'):
            table_buffer.append(line)
            i += 1
            # Collect all table lines
            while i < len(lines) and '|' in lines[i] and lines[i].strip().startswith('|'):
                table_buffer.append(lines[i])
                i += 1
            add_table_from_md(doc, table_buffer)
            table_buffer = []
            continue

        # ── Blockquote ─────────────────────────────────────────
        if line.strip().startswith('>'):
            quote_buffer.append(line)
            i += 1
            while i < len(lines) and lines[i].strip().startswith('>'):
                quote_buffer.append(lines[i])
                i += 1
            text = ' '.join(
                re.sub(r'^>\s*', '', l).strip() for l in quote_buffer
            ).strip()
            # Remove markdown bold/italic
            text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
            text = re.sub(r'\*([^*]+)\*', r'\1', text)
            add_blockquote(doc, text)
            quote_buffer = []
            continue

        # ── Headings ───────────────────────────────────────────
        h4 = re.match(r'^#{4}\s+(.*)', line)
        h3 = re.match(r'^#{3}\s+(.*)', line)
        h2 = re.match(r'^#{2}\s+(.*)', line)
        h1 = re.match(r'^#\s+(.*)', line)

        if h1 and not h2:
            # Skip top-level title (already added)
            i += 1
            continue
        elif h2 and not h3:
            add_heading(doc, h2.group(1).strip(), 1)
        elif h3 and not h4:
            add_heading(doc, h3.group(1).strip(), 2)
        elif h4:
            add_heading(doc, h4.group(1).strip(), 3)

        # ── Horizontal rule ────────────────────────────────────
        elif re.match(r'^---+\s*$', line):
            add_horizontal_rule(doc)

        # ── Bullet list ────────────────────────────────────────
        elif re.match(r'^[-*]\s+', line):
            text = re.sub(r'^[-*]\s+', '', line)
            add_bullet(doc, text, level=0)
        elif re.match(r'^\s{2,}[-*]\s+', line):
            text = re.sub(r'^\s+[-*]\s+', '', line)
            add_bullet(doc, text, level=1)

        # ── Numbered list ──────────────────────────────────────
        elif re.match(r'^\d+\.\s+', line):
            text = re.sub(r'^\d+\.\s+', '', line)
            add_bullet(doc, text, level=0)

        # ── Empty line ─────────────────────────────────────────
        elif line.strip() == '':
            pass  # natural spacing

        # ── Body paragraph ─────────────────────────────────────
        else:
            if line.strip():
                add_body_para(doc, line.strip())

        i += 1

    doc.save(docx_path)
    print(f'Saved: {docx_path}')


# ─────────────────────────────────────────────────────────────
# GENERATE BOTH DOCUMENTS
# ─────────────────────────────────────────────────────────────

md_to_docx(
    md_path=f'{DISSECTION}/UMS_REPORT.md',
    docx_path=f'{DISSECTION}/UMS_REPORT.docx',
    doc_title='LPU University Management System (UMS)',
    doc_subtitle='Technical Architecture Report  |  Paladeium Research  |  2026-09-15'
)

md_to_docx(
    md_path=f'{DISSECTION}/LPU_TOUCH_REPORT.md',
    docx_path=f'{DISSECTION}/LPU_TOUCH_REPORT.docx',
    doc_title='LPU Touch Mobile Application',
    doc_subtitle='Reverse Engineering & Architecture Report  |  Paladeium Research  |  2026-09-16'
)

print('Done.')
